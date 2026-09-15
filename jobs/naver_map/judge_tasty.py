"""posts.json(블로그 본문) → 빵별 '맛있다' 표 수. 규칙 기반, AI 없음.

한 문장 안에 빵 이름 + 긍정 표현이 있고 부정·전해들은 표현이 없으면 그 빵에 1표.
긍정 문장에 빵 이름이 없으면 바로 앞 문장의 빵에 붙인다(블로그는 한 줄씩 끊어 쓰는 경우가 많다).
글 하나에서 같은 빵은 1표. 본문은 화면에 출력하지 않고 엑셀 근거 칸에만 남긴다.

blog.naver.com robots.txt 가 AI 학습·RAG 목적 접근을 금지하므로 본문 판정을 AI 에 맡기지 않는다.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))
from jobs.naver_map.analyze_breads import find_breads, normalize  # noqa: E402

# 공백 제거·소문자화한 문장에 부분 문자열로 맞춘다.
POSITIVE = ("맛있", "맛잇", "맛나", "맛났", "존맛", "꿀맛", "맛도리", "마싯", "마쉿", "마딨", "추천", "강추",
            "ㅊㅊ", "최고", "미쳤", "미침", "짱", "훌륭", "만족", "좋았", "좋더", "재구매", "또사", "또먹",
            "인생빵", "합격", "괜찮았", "괜찮더", "갠츈", "취향저격", "감탄", "jmt")
NEGATIVE = ("별로", "쏘쏘", "soso", "그냥그랬", "그저그랬", "그냥그래", "아쉬", "실망", "맛없", "맛이없",
            "평범", "느끼", "비추", "너무짜", "짰", "안맞", "애매", "밍밍", "퍽퍽", "푸석", "안맛",
            "맛있진않", "맛있지않", "맛있지는않", "맛있는건아니", "그정도는아니", "비해가격")
# 먹기 전 기대·남의 말·겉모습 — 본인이 먹고 한 평가가 아니다.
HEARSAY = ("맛있다고", "맛있대", "맛있다던", "맛있을것", "맛있어보", "맛있게생", "맛나보", "맛있겠",
           "먹어보고싶", "추천받", "추천해준", "추천해줘", "후기보니", "리뷰보니", "맛있다는", "기대")
# 가게·방문을 추천하는 말 — 빵 칭찬이 아니다 (실데이터에서 '카페 추천' 이 가장 흔한 오탐).
PLACE_RECOMMEND = re.compile(r"(카페|이커리|빵집|가게|매장|놀거리|방문|코스|데이트|곳|장소|맛집|시간|오시길)(으로|로)?추천"
                             r"|(맛있|맛나)는.{0,10}?(곳|집|베이커리|카페)")  # '메론빵이 맛있는 베이커리' = 가게 소개
# 앞 문장의 빵에 넘겨도 되는 '맛' 표현만. '좋았다'·'추천'은 분위기·가게에도 쓰여서 넘기지 않는다.
TASTE = ("맛있", "맛잇", "맛나", "맛났", "존맛", "꿀맛", "맛도리", "마싯", "마쉿", "마딨", "미쳤", "미침", "jmt")
# '소보로빵 같은 느낌' 처럼 비교로만 나온 빵 이름은 세지 않는다.
COMPARISON = re.compile(r"(느낌|같|처럼|스타일|st\b|맛이나)")
CARRY_BREAD_MAX, CARRY_TASTE_MAX = 25, 40  # 공백 뺀 글자 수 — 짧은 메뉴 줄 → 짧은 맛 평가만 연결
SPLIT = re.compile(r"(?<=[.!?~♡❤])\s+|[ㅋㅎㅠㅜ]{2,}|\n")


def sentences(paragraphs: list[str]) -> list[str]:
    out = []
    for paragraph in paragraphs:
        out.extend(part.strip() for part in SPLIT.split(paragraph) if part and part.strip())
    return out


def classify(sentence: str) -> str:
    """'neg' | 'hearsay' | 'pos' | '' (평가 없음)."""
    text = PLACE_RECOMMEND.sub("", normalize(sentence))
    if any(word in text for word in NEGATIVE):
        return "neg"
    if any(word in text for word in HEARSAY):
        return "hearsay"
    if any(word in text for word in POSITIVE):
        return "pos"
    return ""


def breads_in(sentence: str) -> list[str]:
    """비교 표현('~같은 느낌') 바로 앞의 빵 이름은 뺀다."""
    from jobs.naver_map.analyze_breads import BREAD_LEXICON
    text = normalize(sentence)
    for aliases in BREAD_LEXICON.values():
        for alias in sorted(aliases, key=len, reverse=True):
            needle = normalize(alias)
            # 별칭이 '소보로' 처럼 '빵' 을 뺀 형태일 수 있어 뒤의 '빵' 도 함께 지운다
            text = re.sub(re.escape(needle) + r"(?:빵)?(?:이랑|이나|과|와|은|는|이|가)?(?=" + COMPARISON.pattern + ")",
                          "", text)
    return list(find_breads(text))


def judge_post(paragraphs: list[str]) -> dict[str, str]:
    """빵 → 근거 문장. 글 하나에서 빵마다 첫 근거만 남긴다."""
    votes: dict[str, str] = {}
    carry: list[str] = []  # 평가 없이 짧게 이름만 나온 직전 빵
    for sentence in sentences(paragraphs):
        breads = breads_in(sentence)
        verdict = classify(sentence)
        length = len(normalize(sentence))
        if verdict == "pos":
            if breads:
                targets, note = breads, ""
            elif length <= CARRY_TASTE_MAX and any(w in normalize(sentence) for w in TASTE):
                targets, note = carry, " (앞 문장 연결)"
            else:
                targets, note = [], ""
            for bread in targets:
                votes.setdefault(bread, sentence[:150] + note)
        if breads:
            carry = breads if not verdict and length <= CARRY_BREAD_MAX else []
        elif verdict:
            carry = []
    return votes


def run(posts: list[dict]) -> tuple[Counter, list[dict]]:
    counts: Counter = Counter()
    rows = []
    for post in posts:
        votes = judge_post(post.get("paragraphs") or [])
        counts.update(votes.keys())
        rows.append({"index": post.get("index"), "blog_url": post.get("blog_url", ""),
                     "status": "본문 없음: " + post["error"] if not post.get("paragraphs") else "",
                     "chars": sum(len(p) for p in post.get("paragraphs") or []), "votes": votes})
    return counts, rows


def write_excel(counts: Counter, rows: list[dict], destination: Path, preview: dict[str, int]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F6D7A")

    def sheet(ws, header, widths):
        ws.append(header)
        for cell, width in zip(ws[1], widths):
            cell.font, cell.fill = header_font, header_fill
            ws.column_dimensions[cell.column_letter].width = width

    ws = wb.active
    ws.title = "본문 기준 순위"
    sheet(ws, ["순위", "빵", "맛있다고 한 글 수(본문·규칙)", "참고: 미리보기·AI 판정"], [6, 20, 26, 22])
    names = sorted(set(counts) | set(preview), key=lambda b: (-counts.get(b, 0), -preview.get(b, 0), b))
    rank, last = 0, None
    for position, bread in enumerate(names, 1):
        if counts.get(bread, 0) != last:
            rank, last = position, counts.get(bread, 0)
        ws.append([rank if counts.get(bread) else "-", bread, counts.get(bread, 0), preview.get(bread, "")])

    ws = wb.create_sheet("글별 판정")
    sheet(ws, ["글 번호", "맛있다고 한 빵", "근거 문장(본문 일부)", "본문 글자 수", "상태", "글 주소"],
          [8, 36, 90, 12, 30, 48])
    for row in rows:
        ws.append([row["index"], ", ".join(row["votes"]) or "(없음)",
                   "\n".join(f"[{b}] {e}" for b, e in row["votes"].items()),
                   row["chars"], row["status"], row["blog_url"]])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("판정 규칙")
    ws.column_dimensions["A"].width = 150
    for line in (
        "판정은 AI 가 아니라 프로그램 규칙으로 했다 (blog.naver.com robots.txt 가 AI 학습·RAG 목적 접근을 금지).",
        "1표 조건: 한 문장에 빵 이름 + 긍정 표현, 부정·전해들은 표현 없음. 긍정 문장에 빵 이름이 없으면 바로 앞 문장의 빵에 연결.",
        "긍정: " + ", ".join(POSITIVE),
        "부정(있으면 제외): " + ", ".join(NEGATIVE),
        "전해들은·기대(있으면 제외): " + ", ".join(HEARSAY),
        "한계: 반어법, 여러 문장 떨어진 평가, '다 맛있었다' 같은 묶음 칭찬, 사전에 없는 빵 이름은 놓치거나 잘못 셀 수 있다. '글별 판정' 근거 문장으로 확인.",
        "'참고: 미리보기·AI 판정' 칸은 목록 미리보기(약 1,000자)를 AI 가 읽고 판정한 이전 결과 — 기준이 달라 숫자가 다를 수 있다.",
    ):
        ws.append([line])
    wb.save(destination)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("posts_json", type=Path, help="fetch_posts.py 가 만든 posts.json")
    parser.add_argument("--output", type=Path, required=True, help="결과 엑셀 경로")
    parser.add_argument("--preview-json", type=Path, help="{빵: 표수} 이전 미리보기 판정 (비교용, 선택)")
    args = parser.parse_args(argv)
    posts = json.loads(args.posts_json.read_text(encoding="utf-8"))
    counts, rows = run(posts)
    preview = json.loads(args.preview_json.read_text(encoding="utf-8")) if args.preview_json else {}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    write_excel(counts, rows, args.output, preview)
    with_votes = sum(bool(r["votes"]) for r in rows)
    missing = sum(bool(r["status"]) for r in rows)
    # 집계 숫자만 출력한다 (본문 문장은 엑셀에만)
    print(f"글 {len(rows)}개 중 본문 없음 {missing}, 1표 이상 {with_votes}")
    for bread, n in counts.most_common():
        print(f"{n:>3}  {bread}")
    print(f"엑셀: {args.output}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
